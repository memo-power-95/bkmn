"""
Top Line Screw Fail - registro diario por turno
================================================
Replica la hoja de Excel 'Top Line Screw Fail':
  - 3 tablas (1Shift, 2Shift, 3Shift) con Screw #1..#4 x (Torque NG, Screw moxibustion)
  - 1 tabla de Total (suma de los 3 turnos)
  - Totales por columna y total general, calculados en vivo

Cada "dia" equivale a una hoja del Excel (09.05, 09.06, 09.15 (2)...).
Todo se guarda automaticamente en SQLite al teclear: si se cierra el programa
no se pierde nada.

Uso:
  - Standalone:   python screw_fail.py
  - Dentro de otra app con ttk.Notebook:
        from screw_fail import ScrewFailTab
        tab = ttk.Frame(notebook); notebook.add(tab, text="Screw Fail")
        ScrewFailTab(tab)
"""

import os
import re
import sys
import sqlite3
import tkinter as tk
from datetime import datetime
from tkinter import ttk, messagebox, filedialog, simpledialog

MAQUINAS = ["Screw #1", "Screw #2", "Screw #3", "Screw #4"]
TURNOS = [1, 2, 3]
CAMPOS = ["torque_ng", "moxibustion"]
ENCABEZADOS = {
    "torque_ng": "Torque NG\n(Screw)",
    "moxibustion": "Screw\nmoxibustion\n(Screw AVI)",
}

VERDE = "#66FF33"
AMARILLO = "#FFFF00"
FUENTE_TITULO = ("Arial", 11, "bold italic")
FUENTE_CELDA = ("Arial", 10, "italic")
FUENTE_NEGRITA = ("Arial", 10, "bold italic")


def _carpeta_base_persistente():
    """Carpeta estable para la BD: junto al .exe si esta empaquetado (PyInstaller),
    o junto al .py si corre como script. Evita que la BD caiga en la carpeta
    temporal que PyInstaller borra al cerrar."""
    if getattr(sys, "frozen", False):
        return os.path.dirname(sys.executable)
    return os.path.dirname(os.path.abspath(__file__))


DB_PATH_DEFAULT = os.path.join(_carpeta_base_persistente(), "screw_fail.db")


# =============================================================================
# Base de datos
# =============================================================================
class ScrewFailDB:
    def __init__(self, path=DB_PATH_DEFAULT):
        self.conn = sqlite3.connect(path)
        self.conn.execute("PRAGMA foreign_keys = ON")
        self._crear_tablas()

    def _crear_tablas(self):
        cur = self.conn.cursor()
        cur.execute("""
            CREATE TABLE IF NOT EXISTS screw_fail_dias (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                etiqueta TEXT UNIQUE NOT NULL,
                creado TEXT
            )
        """)
        cur.execute("""
            CREATE TABLE IF NOT EXISTS screw_fail_datos (
                dia_id INTEGER NOT NULL,
                turno INTEGER NOT NULL,
                maquina TEXT NOT NULL,
                torque_ng INTEGER DEFAULT 0,
                moxibustion INTEGER DEFAULT 0,
                PRIMARY KEY (dia_id, turno, maquina),
                FOREIGN KEY (dia_id) REFERENCES screw_fail_dias(id) ON DELETE CASCADE
            )
        """)
        cur.execute("""
            CREATE TABLE IF NOT EXISTS screw_fail_config (
                clave TEXT PRIMARY KEY,
                valor TEXT
            )
        """)
        self.conn.commit()

    # --- dias ---
    def existe_etiqueta(self, etiqueta):
        cur = self.conn.cursor()
        cur.execute("SELECT 1 FROM screw_fail_dias WHERE etiqueta=?", (etiqueta,))
        return cur.fetchone() is not None

    def etiqueta_disponible(self, base):
        """'09.25' si esta libre; si no '09.25 (2)', '09.25 (3)'... (igual que Excel)."""
        if not self.existe_etiqueta(base):
            return base
        n = 2
        while self.existe_etiqueta(f"{base} ({n})"):
            n += 1
        return f"{base} ({n})"

    def crear_dia(self, etiqueta):
        cur = self.conn.cursor()
        cur.execute(
            "INSERT INTO screw_fail_dias (etiqueta, creado) VALUES (?,?)",
            (etiqueta, datetime.now().isoformat(timespec="seconds")),
        )
        self.conn.commit()
        return cur.lastrowid

    def listar_dias(self):
        cur = self.conn.cursor()
        cur.execute("SELECT id, etiqueta FROM screw_fail_dias ORDER BY id")
        return cur.fetchall()

    def eliminar_dia(self, dia_id):
        cur = self.conn.cursor()
        cur.execute("DELETE FROM screw_fail_datos WHERE dia_id=?", (dia_id,))
        cur.execute("DELETE FROM screw_fail_dias WHERE id=?", (dia_id,))
        self.conn.commit()

    # --- datos ---
    def guardar_valor(self, dia_id, turno, maquina, campo, valor):
        if campo not in CAMPOS:
            raise ValueError(f"Campo invalido: {campo}")
        cur = self.conn.cursor()
        cur.execute(
            f"""INSERT INTO screw_fail_datos (dia_id, turno, maquina, {campo}) VALUES (?,?,?,?)
                ON CONFLICT(dia_id, turno, maquina) DO UPDATE SET {campo}=excluded.{campo}""",
            (dia_id, turno, maquina, int(valor)),
        )
        self.conn.commit()

    def datos_dia(self, dia_id):
        """{(turno, maquina): {'torque_ng': n, 'moxibustion': n}} con ceros por defecto."""
        datos = {(t, m): {c: 0 for c in CAMPOS} for t in TURNOS for m in MAQUINAS}
        cur = self.conn.cursor()
        cur.execute(
            "SELECT turno, maquina, torque_ng, moxibustion FROM screw_fail_datos WHERE dia_id=?",
            (dia_id,),
        )
        for turno, maquina, torque, mox in cur.fetchall():
            if (turno, maquina) in datos:
                datos[(turno, maquina)] = {"torque_ng": torque or 0, "moxibustion": mox or 0}
        return datos

    # --- config ---
    def set_config(self, clave, valor):
        self.conn.execute(
            "INSERT INTO screw_fail_config (clave, valor) VALUES (?,?) "
            "ON CONFLICT(clave) DO UPDATE SET valor=excluded.valor",
            (clave, str(valor)),
        )
        self.conn.commit()

    def get_config(self, clave, default=None):
        cur = self.conn.cursor()
        cur.execute("SELECT valor FROM screw_fail_config WHERE clave=?", (clave,))
        row = cur.fetchone()
        return row[0] if row else default


# =============================================================================
# Exportacion a Excel (mismo acomodo que la hoja original, con formulas)
# =============================================================================
def _nombre_hoja_valido(etiqueta, usados):
    nombre = re.sub(r"[\[\]:*?/\\]", "-", etiqueta)[:31] or "Hoja"
    base, n = nombre, 2
    while nombre in usados:
        sufijo = f" ({n})"
        nombre = base[: 31 - len(sufijo)] + sufijo
        n += 1
    usados.add(nombre)
    return nombre


def escribir_hoja_screw_fail(ws, datos):
    """Escribe una hoja con el mismo layout del Excel original:
    1Shift en D:F, 2Shift en I:K, 3Shift en N:P (filas 4-11) y el Total en D:F (filas 13-20).
    Los totales y la tabla Total son FORMULAS, asi Excel recalcula si alguien edita un numero."""
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import column_index_from_string, get_column_letter

    verde = PatternFill(start_color=VERDE[1:], end_color=VERDE[1:], fill_type="solid")
    amarillo = PatternFill(start_color=AMARILLO[1:], end_color=AMARILLO[1:], fill_type="solid")
    delgado = Side(style="thin", color="000000")
    borde = Border(left=delgado, right=delgado, top=delgado, bottom=delgado)
    centro = Alignment(horizontal="center", vertical="center", wrap_text=True)
    f_titulo = Font(name="Arial", size=11, bold=True, italic=True)
    f_negrita = Font(name="Arial", size=10, bold=True, italic=True)
    f_normal = Font(name="Arial", size=10, italic=True)

    def estilo(celda, fuente, relleno=None):
        celda.font = fuente
        celda.alignment = centro
        celda.border = borde
        if relleno is not None:
            celda.fill = relleno

    def tabla(col_letra, fila0, titulo, valores):
        """valores: lista por maquina de (torque, mox); cada uno int, None o formula."""
        c0 = column_index_from_string(col_letra)
        L = get_column_letter

        t = ws.cell(row=fila0, column=c0, value=titulo)
        t.font = f_titulo
        t.alignment = Alignment(horizontal="center", vertical="center")
        ws.merge_cells(start_row=fila0, start_column=c0, end_row=fila0, end_column=c0 + 2)

        fh = fila0 + 1
        for j, texto in enumerate(["Machine", ENCABEZADOS["torque_ng"], ENCABEZADOS["moxibustion"]]):
            estilo(ws.cell(row=fh, column=c0 + j, value=texto), f_negrita, verde)
        ws.row_dimensions[fh].height = 45

        for i, maquina in enumerate(MAQUINAS):
            fila = fila0 + 2 + i
            estilo(ws.cell(row=fila, column=c0, value=maquina), f_normal)
            torque, mox = valores[i]
            estilo(ws.cell(row=fila, column=c0 + 1, value=torque), f_normal)
            estilo(ws.cell(row=fila, column=c0 + 2, value=mox), f_normal)

        primera, ultima = fila0 + 2, fila0 + 1 + len(MAQUINAS)
        ft = ultima + 1
        estilo(ws.cell(row=ft, column=c0, value="Total"), f_negrita, amarillo)
        estilo(ws.cell(row=ft + 1, column=c0), f_negrita, amarillo)
        ws.merge_cells(start_row=ft, start_column=c0, end_row=ft + 1, end_column=c0)

        for j in (1, 2):
            col = L(c0 + j)
            estilo(ws.cell(row=ft, column=c0 + j, value=f"=SUM({col}{primera}:{col}{ultima})"), f_negrita, amarillo)
        estilo(ws.cell(row=ft + 1, column=c0 + 1, value=f"={L(c0 + 1)}{ft}+{L(c0 + 2)}{ft}"), f_negrita, amarillo)
        estilo(ws.cell(row=ft + 1, column=c0 + 2), f_negrita, amarillo)
        ws.merge_cells(start_row=ft + 1, start_column=c0 + 1, end_row=ft + 1, end_column=c0 + 2)

    columnas_turno = {1: "D", 2: "I", 3: "N"}
    for turno, col in columnas_turno.items():
        valores = []
        for maquina in MAQUINAS:
            d = datos[(turno, maquina)]
            valores.append((d["torque_ng"] or None, d["moxibustion"] or None))
        tabla(col, 4, f"{turno}Shift Top Line Screw Fail", valores)

    # tabla Total: suma por formula de las 3 tablas de turno (fila i de cada una)
    valores_total = []
    for i in range(len(MAQUINAS)):
        fila = 6 + i
        valores_total.append((f"=E{fila}+J{fila}+O{fila}", f"=F{fila}+K{fila}+P{fila}"))
    tabla("D", 13, "Top Line Screw Fail_Total", valores_total)

    for col in "DEFIJKNOP":
        ws.column_dimensions[col].width = 15


def exportar_screw_fail_excel(db, dias, ruta):
    """dias: lista de (dia_id, etiqueta). Una hoja por dia, en el orden recibido."""
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)
    usados = set()
    for dia_id, etiqueta in dias:
        ws = wb.create_sheet(title=_nombre_hoja_valido(etiqueta, usados))
        escribir_hoja_screw_fail(ws, db.datos_dia(dia_id))
    wb.save(ruta)
    return ruta


# =============================================================================
# Interfaz (se puede incrustar en cualquier Frame / pestaña de Notebook)
# =============================================================================
class ScrewFailTab:
    def __init__(self, parent, db=None):
        self.parent = parent
        self.db = db or ScrewFailDB()
        self.dia_actual_id = None
        self.map_etiqueta_id = {}
        self.vars = {}          # (turno, maquina, campo) -> StringVar (editables)
        self.vars_total = {}    # ('total', maquina, campo) -> StringVar (solo lectura)
        self.vars_sumas = {}    # (clave, campo|'gran') -> StringVar
        self.entries = []       # orden de captura para navegar con Enter
        self._cargando = False
        self._construir()
        self._refrescar_dias()

    # ---------------- construccion ----------------
    def _construir(self):
        barra = ttk.Frame(self.parent, padding=8)
        barra.pack(fill="x")

        ttk.Label(barra, text="Día / hoja:").pack(side="left")
        self.combo_dia = ttk.Combobox(barra, width=14, state="readonly")
        self.combo_dia.pack(side="left", padx=5)
        self.combo_dia.bind("<<ComboboxSelected>>", self._on_seleccionar_dia)

        ttk.Button(barra, text="Nuevo día (hoy)", command=self._nuevo_dia_hoy).pack(side="left", padx=4)
        ttk.Button(barra, text="Nuevo con nombre...", command=self._nuevo_dia_con_nombre).pack(side="left", padx=4)
        ttk.Button(barra, text="Eliminar día", command=self._eliminar_dia).pack(side="left", padx=4)
        ttk.Button(barra, text="Exportar este día", command=self._exportar_dia_actual).pack(side="left", padx=(20, 4))
        ttk.Button(barra, text="Exportar todos los días", command=self._exportar_todos).pack(side="left", padx=4)

        self.lbl_estado = ttk.Label(
            self.parent, text="Se guarda automáticamente al escribir. Enter = siguiente celda.",
            foreground="#555555",
        )
        self.lbl_estado.pack(anchor="w", padx=10)

        cuerpo = tk.Frame(self.parent)
        cuerpo.pack(fill="both", expand=True, padx=10, pady=10)

        vcmd = (self.parent.register(self._validar_numero), "%P")
        for turno in TURNOS:
            self._construir_tabla(cuerpo, f"{turno}Shift Top Line Screw Fail", turno, 0, turno - 1, vcmd)
        self._construir_tabla(cuerpo, "Top Line Screw Fail_Total", "total", 1, 0, None)

    def _construir_tabla(self, padre, titulo, clave, fila_grid, col_grid, vcmd):
        marco = tk.Frame(padre)
        marco.grid(row=fila_grid, column=col_grid, padx=18, pady=10, sticky="n")

        tk.Label(marco, text=titulo, font=FUENTE_TITULO).grid(row=0, column=0, columnspan=3, pady=(0, 3))

        celda = dict(relief="solid", bd=1, font=FUENTE_NEGRITA)
        tk.Label(marco, text="Machine", bg=VERDE, width=11, height=3, **celda).grid(row=1, column=0, sticky="nsew")
        for j, campo in enumerate(CAMPOS):
            tk.Label(marco, text=ENCABEZADOS[campo], bg=VERDE, width=12, height=3, **celda).grid(
                row=1, column=j + 1, sticky="nsew"
            )

        for i, maquina in enumerate(MAQUINAS):
            fila = i + 2
            tk.Label(marco, text=maquina, bg="white", relief="solid", bd=1, font=FUENTE_CELDA).grid(
                row=fila, column=0, sticky="nsew"
            )
            for j, campo in enumerate(CAMPOS):
                var = tk.StringVar()
                if clave == "total":
                    self.vars_total[(maquina, campo)] = var
                    tk.Label(marco, textvariable=var, bg="#F2F2F2", relief="solid", bd=1, font=FUENTE_CELDA).grid(
                        row=fila, column=j + 1, sticky="nsew"
                    )
                else:
                    self.vars[(clave, maquina, campo)] = var
                    entry = tk.Entry(
                        marco, textvariable=var, width=10, justify="center", relief="solid", bd=1,
                        font=FUENTE_CELDA, validate="key", validatecommand=vcmd,
                    )
                    entry.grid(row=fila, column=j + 1, sticky="nsew")
                    entry.bind("<Return>", self._siguiente_celda)
                    var.trace_add("write", lambda *_a, t=clave, m=maquina, c=campo: self._on_cambio(t, m, c))
                    self.entries.append(entry)

        f_total = len(MAQUINAS) + 2
        tk.Label(marco, text="Total", bg=AMARILLO, **celda).grid(row=f_total, column=0, rowspan=2, sticky="nsew")
        for j, campo in enumerate(CAMPOS):
            var = tk.StringVar(value="0")
            self.vars_sumas[(clave, campo)] = var
            tk.Label(marco, textvariable=var, bg=AMARILLO, **celda).grid(row=f_total, column=j + 1, sticky="nsew")
        var_gran = tk.StringVar(value="0")
        self.vars_sumas[(clave, "gran")] = var_gran
        tk.Label(marco, textvariable=var_gran, bg=AMARILLO, **celda).grid(
            row=f_total + 1, column=1, columnspan=2, sticky="nsew"
        )

    # ---------------- captura ----------------
    @staticmethod
    def _validar_numero(propuesto):
        return propuesto == "" or (propuesto.isdigit() and len(propuesto) <= 5)

    def _siguiente_celda(self, event):
        idx = self.entries.index(event.widget)
        siguiente = self.entries[(idx + 1) % len(self.entries)]
        siguiente.focus_set()
        siguiente.select_range(0, tk.END)
        return "break"

    def _on_cambio(self, turno, maquina, campo):
        if self._cargando:
            return
        if self.dia_actual_id is None:
            return
        texto = self.vars[(turno, maquina, campo)].get()
        valor = int(texto) if texto.isdigit() else 0
        self.db.guardar_valor(self.dia_actual_id, turno, maquina, campo, valor)
        self._recalcular_totales()

    def _valor(self, turno, maquina, campo):
        texto = self.vars[(turno, maquina, campo)].get()
        return int(texto) if texto.isdigit() else 0

    def _recalcular_totales(self):
        gran_total = {c: 0 for c in CAMPOS}
        for turno in TURNOS:
            for campo in CAMPOS:
                suma = sum(self._valor(turno, m, campo) for m in MAQUINAS)
                self.vars_sumas[(turno, campo)].set(str(suma))
                gran_total[campo] += suma
            self.vars_sumas[(turno, "gran")].set(
                str(sum(int(self.vars_sumas[(turno, c)].get()) for c in CAMPOS))
            )
        for maquina in MAQUINAS:
            for campo in CAMPOS:
                self.vars_total[(maquina, campo)].set(str(sum(self._valor(t, maquina, campo) for t in TURNOS)))
        for campo in CAMPOS:
            self.vars_sumas[("total", campo)].set(str(gran_total[campo]))
        self.vars_sumas[("total", "gran")].set(str(sum(gran_total.values())))

    # ---------------- dias ----------------
    def _refrescar_dias(self, seleccionar_id=None):
        dias = self.db.listar_dias()
        self.map_etiqueta_id = {etq: did for did, etq in dias}
        self.combo_dia["values"] = [etq for _did, etq in dias]

        if seleccionar_id is None:
            guardado = self.db.get_config("ultimo_dia")
            ids = [did for did, _ in dias]
            if guardado and guardado.isdigit() and int(guardado) in ids:
                seleccionar_id = int(guardado)
            elif ids:
                seleccionar_id = ids[-1]

        if seleccionar_id is None:
            self.dia_actual_id = None
            self.combo_dia.set("")
            self._cargar_dia(None)
            return

        for did, etq in dias:
            if did == seleccionar_id:
                self.combo_dia.set(etq)
        self._cargar_dia(seleccionar_id)

    def _on_seleccionar_dia(self, event=None):
        did = self.map_etiqueta_id.get(self.combo_dia.get())
        if did is not None:
            self._cargar_dia(did)

    def _cargar_dia(self, dia_id):
        self.dia_actual_id = dia_id
        self._cargando = True
        try:
            datos = self.db.datos_dia(dia_id) if dia_id is not None else None
            for (turno, maquina, campo), var in self.vars.items():
                valor = datos[(turno, maquina)][campo] if datos else 0
                var.set(str(valor) if valor else "")
        finally:
            self._cargando = False

        estado = "normal" if dia_id is not None else "disabled"
        for entry in self.entries:
            entry.config(state=estado)
        if dia_id is not None:
            self.db.set_config("ultimo_dia", dia_id)
        else:
            self.lbl_estado.config(text="Crea un día con 'Nuevo día (hoy)' para empezar a capturar.")
        self._recalcular_totales()

    def _nuevo_dia_hoy(self):
        etiqueta = self.db.etiqueta_disponible(datetime.now().strftime("%m.%d"))
        nuevo_id = self.db.crear_dia(etiqueta)
        self._refrescar_dias(nuevo_id)
        self.lbl_estado.config(text=f"✔ Día '{etiqueta}' creado. Se guarda automáticamente al escribir.")
        if self.entries:
            self.entries[0].focus_set()

    def _nuevo_dia_con_nombre(self):
        sugerido = self.db.etiqueta_disponible(datetime.now().strftime("%m.%d"))
        etiqueta = simpledialog.askstring(
            "Nuevo día", "Nombre de la hoja (ej. 09.25):", initialvalue=sugerido, parent=self.parent
        )
        if etiqueta is None:
            return
        etiqueta = etiqueta.strip()
        if not etiqueta:
            messagebox.showwarning("Nombre vacío", "Escribe un nombre para el día.", parent=self.parent)
            return
        if re.search(r"[\[\]:*?/\\]", etiqueta) or len(etiqueta) > 31:
            messagebox.showwarning(
                "Nombre inválido",
                "Máximo 31 caracteres y sin estos símbolos: [ ] : * ? / \\ (Excel no los acepta en hojas).",
                parent=self.parent,
            )
            return
        if self.db.existe_etiqueta(etiqueta):
            messagebox.showerror("Ya existe", f"Ya hay un día llamado '{etiqueta}'.", parent=self.parent)
            return
        nuevo_id = self.db.crear_dia(etiqueta)
        self._refrescar_dias(nuevo_id)

    def _eliminar_dia(self):
        if self.dia_actual_id is None:
            return
        etiqueta = self.combo_dia.get()
        if not messagebox.askyesno(
            "Eliminar día",
            f"¿Eliminar el día '{etiqueta}' con todos sus datos?\nNo se puede deshacer.",
            parent=self.parent,
        ):
            return
        self.db.eliminar_dia(self.dia_actual_id)
        self.dia_actual_id = None
        self._refrescar_dias()

    # ---------------- exportacion ----------------
    def _pedir_ruta(self, sugerido):
        return filedialog.asksaveasfilename(
            parent=self.parent, defaultextension=".xlsx", initialfile=sugerido,
            filetypes=[("Excel", "*.xlsx")],
        )

    def _exportar_dia_actual(self):
        if self.dia_actual_id is None:
            messagebox.showinfo("Sin día", "Selecciona o crea un día primero.", parent=self.parent)
            return
        etiqueta = self.combo_dia.get()
        ruta = self._pedir_ruta(f"screw_fail_{etiqueta}.xlsx")
        if ruta:
            self._exportar([(self.dia_actual_id, etiqueta)], ruta)

    def _exportar_todos(self):
        dias = self.db.listar_dias()
        if not dias:
            messagebox.showinfo("Sin datos", "No hay días registrados.", parent=self.parent)
            return
        ruta = self._pedir_ruta(f"screw_fail_{datetime.now():%Y%m%d}.xlsx")
        if ruta:
            self._exportar(dias, ruta)

    def _exportar(self, dias, ruta):
        try:
            exportar_screw_fail_excel(self.db, dias, ruta)
        except PermissionError:
            messagebox.showerror(
                "No se pudo guardar", "El archivo está abierto en Excel. Ciérralo e intenta de nuevo.",
                parent=self.parent,
            )
            return
        messagebox.showinfo("Exportado", f"Archivo guardado:\n{ruta}", parent=self.parent)


if __name__ == "__main__":
    root = tk.Tk()
    root.title("Top Line Screw Fail")
    root.geometry("1150x620")
    notebook = ttk.Notebook(root)
    notebook.pack(fill="both", expand=True)
    tab = ttk.Frame(notebook)
    notebook.add(tab, text="Screw Fail")
    ScrewFailTab(tab)
    root.mainloop()
